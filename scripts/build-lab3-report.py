"""Build from the retained lab 2 DOCX template, using only verified local evidence."""
import argparse,copy,hashlib,importlib.util,json,zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parent.parent
R=ROOT/'tests/performance-tests/results-docker'
B=ROOT/'tests/performance-tests/results'
S=ROOT/'evidence/screenshots/crops'
REF=Path(r'D:/УНИВЕР/МАГІСТРАТУРА/1 СЕМЕСТР/РОЗПОДІЛЕНІ СХОВИЩА ДАНИХ/Лаба 2/TravelerAPI_Lab2_Samonov.docx')
REF_HASH='48192a2d42249ab1455fc3c4fd4086fd5a502ea8be33973bae82e15d4314cc26'
BASE='https://github.com/mishasamonov/traveler-api-dds-lab1'
CODE_COMMIT='c83cc5e5b10bc0207bf19f84228cc29f08c710c4'

spec=importlib.util.spec_from_file_location('lab2_report',ROOT/'scripts/build-lab2-report.py')
helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
spec=importlib.util.spec_from_file_location('comparison',ROOT/'scripts/compare-container-results.py')
comparison=importlib.util.module_from_spec(spec);spec.loader.exec_module(comparison)
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def n(x,d=2):return f'{x:,.{d}f}'.replace(',',' ').replace('.',',')
def count(x):return f'{x:,}'.replace(',',' ')
def pct(x):return n(x*100)+' %'
def ms(x):return n(x)+' мс'

def build(output,partial=False):
 assert hashlib.sha256(REF.read_bytes()).hexdigest()==REF_HASH
 names=[name for name in ('smoke','load','stress','spike','endurance') if (R/f'{name}-summary.json').exists()]
 if not partial:assert len(names)==5 and len(load(R/'endurance-analysis.json')['steady_minutes'])==30
 sums={name:load(R/f'{name}-summary.json')['metrics'] for name in names}
 baselines={name:load(B/f'{name}-summary.json')['metrics'] for name in names}
 analyses={name:load(R/f'{name}-analysis.json') for name in names if name!='smoke'}
 old={name:load(B/f'{name}-analysis.json') for name in analyses}
 d=Document(REF)
 for element in list(d.element.body):
  if element.tag!=qn('w:sectPr'):d.element.body.remove(element)
 figure=0
 def p(text='',style=None):return d.add_paragraph(text,style)
 def h(text):
  a=d.add_heading(text,1);a.paragraph_format.keep_with_next=True
  for run in a.runs:run.bold=True
  return a
 def t(headers,rows,widths):return helper.table(d,headers,rows,widths)
 def link(label,url):
  a=p();helper.hyperlink(a,url,url);a.add_run(' — '+label);a.paragraph_format.keep_together=True
 def fig(path,caption,width=6.85):
  nonlocal figure
  path=Path(path)
  if not path.exists():
   if partial:return
   raise FileNotFoundError(path)
  figure+=1
  # Keep very tall output readable while avoiding an oversized page object.
  with Image.open(path) as im:
   width=min(width,7.7*im.width/im.height)
  a=p();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.keep_with_next=True
  shape=a.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',caption)
  a=p(f'Рисунок {figure}  {caption}','Caption');a.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for run in a.runs:run.bold=True
 def shot(stem,caption):fig(S/(stem+'.png'),caption)
 def metrics(name):
  m=sums[name]
  p(f'За весь запуск виконано {count(m["http_reqs"]["values"]["count"])} запитів; середній час — {ms(m["http_req_duration"]["values"]["avg"])}, p95 — {ms(m["http_req_duration"]["values"]["p(95)"])}. Частка неочікуваних HTTP помилок — {pct(m["http_req_failed"]["values"]["rate"])}.')

 p('Лабораторна робота 3','Title').runs[0].bold=True
 p('Докеризація TravelerAPI та порівняння продуктивності','Subtitle').runs[0].italic=False
 p('Самонов Михайло ІН.м-62/2')
 p('Мета роботи - контейнеризувати TravelerAPI з лабораторної №1 разом із PostgreSQL та оцінити зміну продуктивності за тими самими сценаріями k6, які виконувалися в лабораторній №2.')
 link('Код лабораторної та результати тестів',BASE+'/tree/lab3-docker')
 p('Посилання на коміт із Dockerfile, Compose та засобами перевірки наведене в кінці документа.')

 h('1 Початковий проєкт та середовище')
 p('Використав гілку lab2-performance. REST API, SQL-схему, пакет залежностей і сценарії Hurl та k6 не змінював. Перевірка хешів 24 файлів підтвердила однаковий текст; у дев’яти файлах відрізняються лише переноси рядків Windows і Linux.')
 t(['Компонент','Налаштування'],[
  ['Комп’ютер','Intel Core i5-11300H, 8 логічних процесорів, 15,79 ГБ RAM'],
  ['Середовище Docker','Docker Desktop 4.77.0, Engine 29.5.3, Linux-контейнери на Windows'],
  ['API','Node.js 22.18.0, Express 5, один процес, пул PostgreSQL 20'],
  ['PostgreSQL','16.15, READ COMMITTED, shared_buffers 128 МБ, max_connections 100'],
  ['Ресурси API','2 CPU, 512 MiB RAM; користувач node з UID 1000'],
  ['Ресурси БД','2 CPU, 1 GiB RAM; іменований том postgres_data'],
  ['Генератор','k6 2.3.0 на Windows, звернення до 127.0.0.1:4567'],
 ],[1.55,5.45])
 p('Порівнюю з фактично виміряними версіями з лабораторної №2: Node.js 22.18.0 і PostgreSQL 16.15. Baseline працював нативно у Windows без контейнерних лімітів. У новому запуску API та БД працюють у Linux-середовищі Docker Desktop, тому різниця включає вплив ОС, віртуалізації, мережі, дискового вводу-виводу та встановлених лімітів.')

 h('2 Контейнеризація та запуск')
 p('Додав багатостадійний Dockerfile. На першій стадії npm ci --omit=dev встановлює залежності з lock-файла. У фінальний образ копіюються тільки потрібні файли. Процес працює від користувача node; файлова система API доступна лише для читання, а /tmp розміщено у tmpfs.')
 p('У docker-compose.yml описані два сервіси: postgres і web. Вони спілкуються в bridge-мережі backend. На комп’ютер опублікований тільки порт API 4567. PostgreSQL зберігає дані в іменованому томі; web очікує успішний healthcheck БД перед застосуванням схеми та запуском сервера.')
 p('Пароль і DATABASE_URL читаються з .env. Цей файл внесено до .gitignore та .dockerignore. У репозиторії збережено .env.example з безпечними прикладами. Для обох сервісів налаштовані healthcheck, ліміти CPU і пам’яті, restart та ротація журналів. PostgreSQL після ініціалізації виконується від користувача postgres.')
 p('Створив .env за прикладом та послідовно виконав команди:')
 for line in ['docker --version','docker compose build','docker compose up -d --wait']:p(line)
 p('Повний проєкт також запускається командою docker compose up --build після створення .env. Окремий docker-compose.dev.yml додає монтування коду, node --watch і окремий том БД для розроблення. Для оновлення образів передбачені docker compose pull та docker compose build --pull.')
 shot('01-docker-desktop','Сервіси PostgreSQL і TravelerAPI у Docker Desktop')
 p('Перевірив http://localhost:4567/health. Відповідь підтверджує доступність API та БД.')
 shot('16-health','Відповідь перевірки здоров’я API')

 h('3 Моніторинг контейнерів')
 p('Виконав docker compose ps, docker compose logs -f web та docker stats. Нижче показані справжні знімки збереженого виводу команд у локальному переглядачі. Вихідні текстові журнали збережені поряд із зображеннями; для читабельності довгі рядки переносяться без зміни чисел.')
 shot('02-compose-ps','Обидва сервіси працюють та мають стан healthy')
 shot('03-web-logs','Запуск API та таймаут пулу з журналу web')
 p('Повідомлення timeout exceeded when trying to connect з’явилося під час перевантаження. Воно пов’язане з очікуванням вільного з’єднання у пулі, для якого в API встановлено таймаут 5 секунд.')
 shot('04-docker-stats','Знімок CPU і пам’яті під навантаженням')
 p('На цьому знімку PostgreSQL використовує близько 192 % CPU. У docker stats 100 % відповідає одному логічному процесору, тому для сервісу з лімітом 2 CPU значення може наближатися до 200 %. Це узгоджується з наближенням БД до встановленого ліміту; окреме профілювання причин затримок не виконувалося.')

 h('4 Функціональна перевірка Hurl')
 p('Запустив оригінальні файли crud.hurl, validation.hurl, management.hurl та race-conditions.hurl через Hurl 8.0.0 у мережі контейнерів. Клієнт звертався до http://web:4567. Виконано 56 запитів у чотирьох файлах, невдалих файлів немає.')
 shot('05-hurl','Успішне виконання оригінальних сценаріїв Hurl')
 p('Додатково пройшли всі три наявні тести конкурентного доступу. Вони перевіряють конфлікт версій плану, унікальний порядок паралельно доданих локацій та повторне оновлення локації після 409. Тестовий контейнер використовував мережевий простір API, щоб незмінений localhost у тесті адресував саме web.')
 shot('06-concurrency','Три успішні тести конкурентного доступу')

 h('5 Методика порівняння з лабораторною 2')
 p('Перед кожним сценарієм відновлював той самий набір: 2 000 планів і по три локації на план. Дані готуються поза вимірюваним інтервалом. Кожен VU працює зі своїм планом; типовий цикл містить вісім читань, два оновлення з version та п’ять пауз по 100 мс. Кожен десятий цикл створює і видаляє тимчасовий план із локацією.')
 t(['Тест','Незмінений профіль'],[
  ['Smoke','20 с, максимум 5 VU'],['Load','25, 50, 100 і 200 VU; плато по 30 с'],
  ['Stress','200, 500, 1 000 і 2 000 VU; плато по 45 с; потім 20 VU'],
  ['Spike','20 до 1 000 VU за 1 с; пік 30 с; спад до 20 VU за 1 с'],
  ['Endurance','25 VU протягом 1 800 с; по 30 с розгін і завершення'],
 ],[1.35,5.65])
 p('Сценарії виконував послідовно, по одному запуску. Критерій прийнятної роботи: p95 < 500 мс і http_req_failed < 1 %. p95 є 95-м перцентилем, а не середнім часом. Окремо збережені середній час, кількість запитів і CPU та пам’ять контейнерів. Очікувані 409 і 404 у smoke не вважаються помилками.')
 p('Для сталих плато відкидаю перші та останні 5 секунд. Під час відновлення також відкидаю 30 секунд gracefulRampDown. Підсумок усього запуску й метрики сталого плато описують різні інтервали, тому в таблицях їх наведено окремо. Закрита модель VU зменшує фактичний потік запитів при збільшенні затримки.')

 h('6 Підсумки однакових сценаріїв')
 rows=[]
 for name in names:
  a=baselines[name]['http_req_duration']['values']['p(95)'];b=sums[name]['http_req_duration']['values']['p(95)']
  rows.append([name.capitalize(),n(a),n(b),('+' if b>=a else '')+n((b/a-1)*100,1)+' %',pct(sums[name]['http_req_failed']['values']['rate'])])
 t(['Тест','p95 Лаб 2\nмс','p95 Docker\nмс','Зміна p95','Помилки\nDocker'],rows,[1.05,1.35,1.35,1.65,1.6])
 p('У всіх п’яти baseline-запусках лабораторної №2 частка HTTP помилок дорівнювала 0 %. Нові stress і spike порушили пороги і затримки, і помилок. Код завершення k6 99 означає порушення порогів; він сам по собі не означає зупинку контейнера.')
 metrics('smoke')
 p(f'Smoke пройшов {count(sums["smoke"]["checks"]["values"]["passes"])} перевірки без збоїв; очікуваних конфліктів версій — {count(sums["smoke"]["expected_version_conflicts"]["values"]["count"])}.')
 shot('07-smoke','Підсумок smoke тесту в Docker')

 h('7 Навантажувальне тестування Load')
 rows=[]
 for key,a in old['load']['stable_plateaus'].items():
  b=analyses['load']['stable_plateaus'][key]
  rows.append([str(a['target_vus']),n(a['p95_ms']),n(b['p95_ms']),n(a['rps'],1),n(b['rps'],1)])
 t(['VU','p95 Лаб 2\nмс','p95 Docker\nмс','Запитів за с\nЛаб 2','Запитів за с\nDocker'],rows,[.6,1.5,1.5,1.7,1.7])
 top=analyses['load']['stable_plateaus']['load200']
 p(f'При 100 VU Docker ще вкладається в поріг p95 500 мс. При 200 VU p95 становить {ms(top["p95_ms"])}, а потік — {n(top["rps"],1)} запитів за секунду. HTTP помилок немає. Найбільше успішно перевірене навантаження в цьому профілі зменшилося з 200 VU у baseline до 100 VU у Docker; точну межу між ними не шукав.')
 fig(R/'load-comparison.png','Затримки Load та частка помилок у Docker')
 shot('08-load','Підсумок Load з перевищенням порога p95')

 h('8 Стресове тестування Stress')
 rows=[]
 for key,b in analyses['stress']['stable_plateaus'].items():
  a=old['stress']['stable_plateaus'][key]
  rows.append(['Відновлення' if key=='recovery' else str(b['target_vus']),n(a['p95_ms']),n(b['p95_ms']),pct(b['failure_rate'])])
 t(['VU або фаза','p95 Лаб 2\nмс','p95 Docker\nмс','Помилки\nDocker'],rows,[1.6,1.8,1.8,1.8])
 metrics('stress')
 p('Перший перевірений деградований рівень у Docker — 200 VU. Під час найвищого навантаження з’являються серверні 500 через таймаут очікування з’єднання з PostgreSQL. Контейнери залишилися запущеними: перезапусків і OOM не зафіксовано.')
 fig(R/'stress-comparison.png','Stress та помилки під час перевантаження і спаду')
 shot('09-stress','Порушені пороги Stress у збереженому виводі k6')
 p('У сирих метриках stress є 3 635 відповідей 500 і 5 989 відповідей 400. Частина 400 виникає через слабке місце незміненого клієнта: JSON помилки може бути прийнятий за поточний план, після чого формується URL з id=undefined. Тому весь відсоток помилок не можна приписувати лише серверу або Docker.')
 p('На сталому відрізку відновлення p95 зменшується до 20,28 мс, але помилки ще становлять 2,54 %. Повного відновлення за спільним критерієм p95 < 500 мс та помилки < 1 % у межах цього сценарію не підтверджено. Позначка PASS для recovery у k6 перевіряє лише затримку.')
 p('Після stress у БД залишилися 2 034 плани та 6 034 локації, з них 34 тимчасові плани тесту; локацій без батьківського плану немає. Це означає незавершене тестове прибирання після невдалих запитів. Перед наступним сценарієм набір відновлено.')

 h('9 Тестування сплеску Spike')
 metrics('spike')
 rows=[]
 for key,label in [('before','До сплеску'),('spike','Сплеск'),('recovery','Відновлення')]:
  a=old['spike']['stable_plateaus'][key];b=analyses['spike']['stable_plateaus'][key]
  rows.append([label,n(a['p95_ms']),n(b['p95_ms']),pct(b['failure_rate'])])
 t(['Фаза','p95 Лаб 2\nмс','p95 Docker\nмс','Помилки\nDocker'],rows,[1.6,1.8,1.8,1.8])
 fig(R/'spike-comparison.png','Spike та збереження помилок клієнта після сплеску')
 shot('10-spike','Підсумок Spike з помилками та перевищенням затримки')
 p('У spike зафіксовано 1 084 відповіді 500 та 3 281 відповідь 400. Після спаду p95 повернувся нижче 500 мс, однак у сталому фрагменті відновлення помилок 15,64 %, переважно через id=undefined. Тому повне відновлення клієнтського сценарію не підтверджено. Після spike перевірка БД показала 2 000 планів, 6 000 локацій і відсутність тимчасових та осиротілих записів.')

 h('10 Тривале тестування Endurance')
 p('Профіль залишив без скорочення: 25 VU протягом 30 хвилин, а з розгоном і завершенням — 31 хвилина. Новий тест почався з відновленого набору даних. Перші та останні п’ять хвилин сталого навантаження порівнюю без розгону й завершення.')
 shot('11-endurance-start','Початок повного 31-хвилинного профілю Endurance')
 if 'endurance' in analyses:
  e=analyses['endurance'];metrics('endurance')
  t(['Інтервал','p95 Лаб 2\nмс','p95 Docker\nмс','Помилки\nDocker'],[
   ['Перші 5 хв',n(old['endurance']['first_5min']['p95_ms']),n(e['first_5min']['p95_ms']),pct(e['first_5min']['failure_rate'])],
   ['Останні 5 хв',n(old['endurance']['last_5min']['p95_ms']),n(e['last_5min']['p95_ms']),pct(e['last_5min']['failure_rate'])],
  ],[1.6,1.8,1.8,1.8])
  minute_p95=[x['p95_ms'] for x in e['steady_minutes'].values()]
  p(f'Хвилинні p95 протягом сталого інтервалу: від {ms(min(minute_p95))} до {ms(max(minute_p95))}. Зміна p95 останніх п’яти хвилин відносно перших — {n((e["last_5min"]["p95_ms"]/e["first_5min"]["p95_ms"]-1)*100,1)} %. Ці дані описують один 30-хвилинний інтервал і не доводять відсутність витоків пам’яті за тривалішої роботи.')
  p('Усі 30 хвилин залишалися в межах SLO 500 мс, хоча наприкінці затримка збільшилася. Причину цієї зміни одним запуском не ізольовано; для перевірки поступового погіршення потрібні повторні вимірювання та профілювання.')
  fig(R/'endurance-comparison.png','Порівняння Endurance протягом повного запуску')
  shot('13-endurance','Підсумкові метрики Endurance')
  shot('12-endurance-finish','Завершення Endurance та кількість повних ітерацій')

 h('11 Споживання ресурсів')
 rows=[]
 for name in names:
  a=comparison.stats(name)
  rows.append([name.capitalize(),n(a['web']['cpu_peak_pct'],1),n(a['postgres']['cpu_peak_pct'],1),n(a['web']['memory_peak_mib'],1),n(a['postgres']['memory_peak_mib'],1)])
 t(['Тест','CPU web\nмакс %','CPU БД\nмакс %','RAM web\nмакс MiB','RAM БД\nмакс MiB'],rows,[1.05,1.35,1.35,1.65,1.6])
 p('Це максимуми серед збережених docker stats з інтервалом приблизно 7 секунд з урахуванням часу виконання команди. Коротші піки могли залишитися між вимірюваннями. CPU наведений у відсотках одного ядра. Пам’ять Docker не зіставляю безпосередньо з приватною пам’яттю процесу Windows із лабораторної №2, оскільки спосіб обліку відрізняється. Витрати всієї VM Docker Desktop до суми цих двох контейнерів не включені.')

 h('12 Збереження даних та автоматична перевірка')
 p('Окремо після навантажувальних тестів перевірив іменований том: додав контрольний план, виконав docker compose down без -v, знову підняв сервіси та прочитав той самий ідентифікатор. Після перевірки видалив тільки контрольний запис. Результат збережений у журналі persistence.')
 shot('15-persistence','Контрольний запис зберігся після перестворення контейнерів')
 p('Додав workflow Docker laboratory 3. На Ubuntu runner він збирає та запускає обидва контейнери, перевіряє /health і UID 1000, виконує Hurl, конкурентні тести та k6 smoke і зберігає діагностику як artifact. Усі кроки на коміті з Dockerfile пройшли. Продуктивність CI з локальними вимірюваннями не порівнюю.')
 shot('14-ci','Успішна перевірка Docker у GitHub Actions')

 h('13 Відтворення та висновки')
 p('Для повторення потрібно створити .env, запустити docker compose up --build та перевірити /health. Функціональні команди наведено в docs/lab3-docker.md. Для вимірювань у виділеній БД traveler_lab2 слід увімкнути LAB3_ALLOW_RESET=true та виконати scripts/run-container-performance.py з параметром --k6 і шляхом до k6 2.3.0. Скрипт виконує п’ять сценаріїв послідовно. Сирі JSON, HTML dashboard, підсумки, часові мітки та docker stats збережені в results-docker; baseline залишений окремо в results.')
 p('Контейнеризація забезпечила запуск API та PostgreSQL однією командою і збереження даних після перестворення контейнерів. Функціональні та конкурентні перевірки пройшли. У виміряній конфігурації найбільше перевірене навантаження Load, яке відповідає SLO, становило 100 VU проти 200 VU у baseline. Stress і spike виявили затримки, таймаути пулу та помилки клієнтського сценарію після серверних збоїв.')
 if 'endurance' in sums:
  m=sums['endurance'];p(f'За повний Endurance при 25 VU p95 становить {ms(m["http_req_duration"]["values"]["p(95)"])}, частка HTTP помилок — {pct(m["http_req_failed"]["values"]["rate"])}. Висновок стосується цієї конфігурації та цього інтервалу вимірювання.')
 p('Отримана різниця є результатом переходу з нативних процесів Windows до Linux-контейнерів Docker Desktop з обмеженнями ресурсів. Один історичний baseline і один новий запуск кожного профілю не ізолюють власну накладну вартість Docker та не дають універсальної оцінки для інших комп’ютерів.')
 h('Посилання')
 link('Гілка лабораторної №3',BASE+'/tree/lab3-docker')
 link('Коміт із Dockerfile, Compose та перевірками',BASE+'/commit/'+CODE_COMMIT)
 if (ROOT/'evidence/commits.json').exists():link('Коміт із вимірюваннями та доказами',BASE+'/commit/'+load(ROOT/'evidence/commits.json')['results'])
 link('Успішний запуск Docker у GitHub Actions',BASE+'/actions/runs/40739221001')
 link('Baseline лабораторної №2',BASE+'/tree/lab2-performance')
 p('Методичні вказівки: SumDU.CS.DDS.Lab3.pdf. Додатково використані офіційні описи Docker Compose та встановлення Hurl.')
 link('Порядок запуску сервісів Docker Compose','https://docs.docker.com/compose/how-tos/startup-order/')
 link('Офіційні способи запуску Hurl','https://hurl.dev/docs/installation.html')

 # Direct typography overrides preserve the retained styles and footer package parts.
 for run in d.element.xpath('.//w:r'):
  props=run.find(qn('w:rPr'))
  if props is None:props=OxmlElement('w:rPr');run.insert(0,props)
  font=props.find(qn('w:rFonts'))
  if font is None:font=OxmlElement('w:rFonts');props.insert(0,font)
  font.attrib.clear()
  for key in ('ascii','hAnsi','eastAsia','cs'):font.set(qn('w:'+key),'Times New Roman')
  for tag in ('sz','szCs'):
   size=props.find(qn('w:'+tag))
   if size is None:size=OxmlElement('w:'+tag);props.append(size)
   size.set(qn('w:val'),'28')
 d.core_properties.title='Лабораторна робота 3 Докеризація TravelerAPI'
 d.core_properties.subject='Розподілені сховища даних'
 output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);d.save(output)
 with zipfile.ZipFile(REF) as a,zipfile.ZipFile(output) as b:
  preserve=[x for x in a.namelist() if x in ('word/styles.xml','word/stylesWithEffects.xml','word/numbering.xml','word/fontTable.xml') or x.startswith(('word/theme/','word/footer','word/header','customXml/'))]
  # python-docx may reserialize opaque parts. Restore the unchanged originals.
  temp=output.with_suffix('.preserve.docx')
  with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
   for info in b.infolist():z.writestr(info,a.read(info.filename) if info.filename in preserve else b.read(info.filename))
 temp.replace(output)
 with zipfile.ZipFile(REF) as a,zipfile.ZipFile(output) as b:
  assert all(a.read(x)==b.read(x) for x in preserve)
 print(json.dumps({'path':str(output),'figures':figure,'complete':not partial,'preserved_parts':len(preserve)},ensure_ascii=False))

if __name__=='__main__':
 a=argparse.ArgumentParser();a.add_argument('--output',required=True);a.add_argument('--partial',action='store_true');x=a.parse_args();build(x.output,x.partial)
