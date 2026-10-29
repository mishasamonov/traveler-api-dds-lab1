"""Create the Ukrainian report using the retained laboratory 2 Word template."""
from pathlib import Path
import copy,hashlib,importlib.util,json,zipfile
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
ROOT=Path(__file__).resolve().parents[2];E=ROOT/'evidence/lab4';S=E/'screenshots/crops'
REF=Path(r'D:/УНИВЕР/МАГІСТРАТУРА/1 СЕМЕСТР/РОЗПОДІЛЕНІ СХОВИЩА ДАНИХ/Лаба 2/TravelerAPI_Lab2_Samonov.docx')
REF_HASH='48192a2d42249ab1455fc3c4fd4086fd5a502ea8be33973bae82e15d4314cc26'
BASE='https://github.com/mishasamonov/traveler-api-dds-lab1'
def load(p):return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def number(v,d=2):return f'{v:,.{d}f}'.replace(',',' ').replace('.',',')
def pct(v):return number(v*100)+' %'
def ms(v):return number(v)+' мс'
def build(output):
 assert hashlib.sha256(REF.read_bytes()).hexdigest()==REF_HASH
 a=load(E/'analysis.json');m=a['load'];p=a['plateaus'];pub=load(E/'commits.json')
 spec=importlib.util.spec_from_file_location('lab2',ROOT/'scripts/build-lab2-report.py');helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
 d=Document(REF)
 for x in list(d.element.body):
  if x.tag!=qn('w:sectPr'):d.element.body.remove(x)
 figure=0
 def text(s='',style=None):return d.add_paragraph(s,style)
 def heading(s):
  par=d.add_heading(s,1);par.paragraph_format.keep_with_next=True
  return par
 def table(headers,rows,widths):return helper.table(d,headers,rows,widths)
 def link(label,url):
  par=text();par.paragraph_format.keep_together=True
  el=OxmlElement('w:hyperlink');el.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True))
  run=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');pr.append(c)
  u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);run.append(pr);t=OxmlElement('w:t');t.text=url;run.append(t);el.append(run);par._p.append(el)
  par.add_run(' — '+label)
 def image(path,caption):
  nonlocal figure
  path=Path(path);assert path.exists(),path;figure+=1
  im=Image.open(path);width=min(6.85,6.6*im.width/im.height)
  par=text();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.keep_with_next=True
  shape=par.add_run().add_picture(str(path),width=Inches(width));shape._inline.docPr.set('descr',caption)
  cap=text(f'Рисунок {figure}  {caption}','Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for run in cap.runs:run.bold=True
 def shot(name,caption):image(S/(name+'.png'),caption)
 names={'baseline':'Без репліки','logical':'Логічна','physical':'Фізична асинхронна','sync':'Фізична синхронна','pgcat':'PgCat і дві репліки'}

 text('Лабораторна робота 4','Title')
 text('Реплікація PostgreSQL у TravelerAPI','Subtitle')
 text('Самонов Михайло ІН.м-62/2')
 text('Мета роботи - налаштувати реплікацію PostgreSQL у контейнеризованому TravelerAPI, перевірити однаковість даних та порівняти швидкодію додатка до і після введення реплік.')
 text('Налаштував логічну й фізичну реплікацію, синхронне підтвердження записів, дві репліки та PgCat. Перевірив роботу додатка тестами Hurl і порівняв швидкодію за сценаріями k6.')

 heading('1 Початковий проєкт і середовище')
 text('Продовжив TravelerAPI з гілки lab3-docker. Реалізацію REST API, початкові сценарії Hurl, smoke і змішаний Load із лабораторної №2 залишив без змін. Перевірка байтів і SHA-256 для 11 ключових файлів підтвердила їхню однаковість із лабораторною №3. Новий read-test.js використовується лише для окремого досліду читання через PgCat.')
 table(['Компонент','Налаштування'],[
  ['Комп’ютер','Intel Core i5-11300H, 8 логічних процесорів, 15,79 ГБ RAM'],
  ['API','Node.js 22.18.0, Express 5, один процес, пул 20 з’єднань'],
  ['База даних','PostgreSQL 16.15 alpine; ізольовані іменовані томи'],
  ['Ресурси','API 2 CPU і 512 MiB; кожна БД 2 CPU і 1 GiB; PgCat 1 CPU і 256 MiB'],
  ['Інструменти','Docker Desktop, Hurl 8.0.0, k6 2.3.0 на Windows'],
 ],[1.5,5.5])
 text('Створив окремі Compose-проєкти traveler_lab4_baseline, logical, physical, sync, logical2, two та pgcat. Старі томи лабораторних не видаляв. На нових томах PostgreSQL виконує ініціалізацію з нуля. Сценарії запускав послідовно; одночасного навантаження різних топологій немає.')

 heading('2 Налаштування логічної реплікації')
 text('Зберіг структуру db/migrations і db/replications/common, publisher, subscriber. Сценарій 500_orchestrate.sh застосовує SQL-схему та сценарії відповідного вузла. На publisher створюються роль реплікації і PUBLICATION tp_pub FOR ALL TABLES. На subscriber створюється така сама схема та SUBSCRIPTION tp_sub1 із copy_data=true й окремим слотом.')
 text('На primary задано listen_addresses=*, wal_level=logical, max_replication_slots=10, max_wal_senders=10, max_logical_replication_workers=4 і max_worker_processes=10. Subscriber очікує готовність primary. Для другого логічного subscriber передбачені інше ім’я підписки, інший слот та окремий том.')
 text('Команди початкового запуску:')
 text('python scripts/lab4/bootstrap-env.py')
 text('python scripts/lab4/lab.py up logical')
 text('Перший скрипт генерує незалежні випадкові паролі у локальному .env. Другий запускає відповідний профіль Compose. Для однієї логічної репліки після підготовки .env також працює docker compose up --build.')
 shot('03-logical-subscriptions','Підписка активна і початкова синхронізація обох таблиць завершена')
 text('Значення srsubstate=r означає готовність таблиці до звичайної реплікації. У pg_stat_replication підключення tp_sub1 має стан streaming та sync_state=async.')

 heading('3 Однаковість даних і функціональні перевірки')
 text('Перед кожним навантаженням відновлював 2 000 планів і 6 000 локацій за тим самим SQL, що в лабораторній №2. Порівнював кількість записів і MD5 впорядкованого JSON усіх рядків кожної таблиці. Це перевіряє ідентичність повного вмісту, а не лише кількості записів. MD5 тут використовується як контрольна сума для порівняння даних.')
 shot('02-logical-equality','Publisher і subscriber містять однакові рядки та контрольні суми')
 text('Окремо створив контрольний план, змінив його назву і версію та видалив його. Після кожної з трьох операцій очікував збігу повних контрольних сум на всіх вузлах. Усі INSERT, UPDATE і DELETE перенесені. Після кожного завершеного тесту збіг даних перевіряв повторно.')
 text('Оригінальні Hurl-сценарії перевіряють CRUD, валідацію, керування локаціями та конфлікти версій. Для логічної топології виконано 56 запитів у чотирьох файлах, невдалих файлів немає. Smoke також проходить усі пороги; очікувані 409 і 404 визначені тестом та не вважаються неочікуваними HTTP помилками.')
 shot('04-logical-smoke','Успішний smoke k6 для логічної реплікації')
 text('Логічна реплікація не переносить DDL. Після CREATE TABLE lab4_ddl_probe на publisher таблиця була відсутня на subscriber. Після застосування DDL вручну та ALTER SUBSCRIPTION REFRESH PUBLICATION рядок із цієї таблиці успішно перенісся.')
 shot('05-logical-ddl','Дослід із DDL на логічній репліці')

 heading('4 Методика порівняння швидкодії')
 text('Для основного порівняння вибрав незмінений Load із лабораторної №2. Профіль триває 170 секунд: розгін до 25, 50, 100 і 200 VU, кожне плато по 30 секунд, переходи по 10 секунд. Звичайний цикл містить вісім читань, два оновлення з version і п’ять пауз по 100 мс; кожен десятий цикл також створює й видаляє тимчасовий план із локацією.')
 text('Критерії: p95 < 500 мс, http_req_failed < 1 %, checks > 99 %. У підсумковій таблиці використані метрики всього запуску. Для порівняння сталих плато відкидав по п’ять секунд з кожного краю: інтервали 15-35, 55-75, 95-115 та 135-155 секунд від початку сценарію. Значення обчислені за фактичним часом завершення HTTP-запитів із сирого JSON k6.')
 text('Основні п’ять конфігурацій виміряні по одному разу з однаковим набором даних та лімітами primary й API. Для всіх залишено wal_level=logical, щоб додатково не змінювати обсяг WAL через інший його рівень. У фізичній конфігурації достатньо wal_level=replica, але тут вищий рівень збережено як контрольну умову.')
 text('Закрита модель VU знижує фактичний потік запитів, коли затримка зростає. Усі процеси працюють на одному ноутбуці, тому малі відмінності можуть залежати від фонових процесів, кешу та планувальника. Три повтори зроблено окремо для досліду читання через PgCat.')
 table(['Топологія','Середній час','p95','Запитів за с','HTTP помилки'],[[names[k],ms(v['avg_ms']),ms(v['p95_ms']),number(v['rps'],1),pct(v['failure_rate'])] for k,v in m.items()],[2.0,1.25,1.25,1.25,1.25])
 text('Код завершення k6 99 означає порушення хоча б одного порога. Він не означає аварійного завершення API. Статус кожного порога збережений у summary.json і summary.txt, а всі вимірювання - у файлах *-raw.json.gz.')
 shot('06-baseline-load','Початковий Load без активної реплікації')
 shot('07-logical-load','Той самий Load із логічною реплікою')
 change=100*(m['logical']['p95_ms']/m['baseline']['p95_ms']-1)
 text(f'Порівняно з baseline p95 логічного варіанта змінився на {number(change,1)} %. У цьому запуску додавання логічної копії не прискорило API: читання й надалі спрямовані на primary, а реплікація додає передавання й застосування змін. Ця невелика різниця одного запуску не є статистичним доказом окремої накладної вартості реплікації.')

 heading('5 Фізична асинхронна реплікація')
 text('Замість незалежної ініціалізації SQL-схеми фізична репліка отримує копію всього кластера через pg_basebackup -R -X stream -C -S replica1 --checkpoint=fast. Параметр -R створює standby.signal і параметри підключення; окремий слот утримує потрібний WAL. Пароль зберігається у .pgpass усередині тому з правами 0600, каталог PGDATA має права 0700. Standby працює з hot_standby=on, а max_worker_processes узгоджений із primary.')
 text('На фізичній репліці не створюю публікацію, підписку або таблиці вручну. Вона відтворює WAL усього кластера та залишається доступною для читання. Для надійного одночасного створення двох копій додав wal_keep_size=256MB і повторне використання вже створеного фізичного слота після перерваного pg_basebackup. Однаковість таблиць і перенесення контрольних INSERT, UPDATE та DELETE підтверджені тими самими перевірками.')
 shot('08-physical-equality','Однакові дані primary та фізичної репліки')
 shot('09-physical-ddl','Нова таблиця й рядок автоматично з’явилися на фізичній репліці')
 text(f'Той самий Load дав середній час {ms(m["physical"]["avg_ms"])}, p95 {ms(m["physical"]["p95_ms"])} і {number(m["physical"]["rps"],1)} запитів за секунду. Для порівняння логічний варіант мав p95 {ms(m["logical"]["p95_ms"])}. Відмінність між ними в цих одиничних вимірюваннях невелика; перевага фізичної схеми в цьому досліді - автоматичне перенесення DDL.')
 shot('11-physical-load','Підсумок Load для фізичної асинхронної реплікації')

 heading('6 Синхронна фізична реплікація')
 text('Після готовності standby встановив synchronous_standby_names=FIRST 1 (replica1) та synchronous_commit=remote_apply, виконав pg_reload_conf(). Налаштування застосовую після pg_basebackup, щоб ініціалізація не чекала репліки, якої ще немає. Значення remote_apply чекає не лише збереження WAL, а й застосування транзакції на standby, тому підтверджений запис уже видимий для її читачів.')
 text('Перевірив фактичну поведінку: призупинив контейнер replica1, виконав UPDATE на primary і спостерігав wait_event=SyncRep. Запит залишався незавершеним до відновлення репліки, після unpause завершився UPDATE 1. Це підтверджує роботу синхронного підтвердження запису.')
 shot('12-sync-block','Commit очікує SyncRep до відновлення репліки')
 text(f'У Load синхронний варіант має p95 {ms(m["sync"]["p95_ms"])} проти {ms(m["physical"]["p95_ms"])} в асинхронному та потік {number(m["sync"]["rps"],1)} запитів за секунду. Очікування remote_apply додає залежність часу commit від мережі й застосування WAL; вплив на змішаний HTTP-сценарій оцінюю за збереженими цифрами, без припущення про універсальне прискорення чи уповільнення.')
 shot('13-sync-load','Load із синхронною фізичною реплікацією')
 text('Якщо обов’язкова синхронна репліка недоступна, записи чекають. Ця конфігурація підвищує вимоги до узгодженості ціною доступності запису.')

 heading('7 Облікові дані та мережеві обмеження')
 text('У SQL і shell-сценаріях немає явно записаного пароля для реплікації. REPL_USER і REPL_PASSWORD читаються із середовища та передаються до psql через змінні. Форматування SQL-ідентифікаторів і значень виконується format із %I та %L. Паролі .env генеруються локально; .env і .runtime виключені з Git, контексту Docker та архіву проєкту.')
 shot('10-security-config','Роль реплікації створюється зі змінних оточення')
 text('Розділив адміністратора, користувача API та користувача реплікації. API має SELECT, INSERT, UPDATE і DELETE, але не має SUPERUSER, REPLICATION, CREATE ROLE, CREATE DATABASE або TRUNCATE. Роль реплікації має LOGIN, REPLICATION і читання таблиць. На логічному subscriber для ролі API встановлений default_transaction_read_only=on.')
 text('Порти PostgreSQL і PgCat не публікуються на хості. Єдиний доступний із хоста порт - API 127.0.0.1:4567. Backend-мережа має internal=true; лише API додатково входить у frontend-мережу. listen_addresses=* потрібен для зв’язку контейнерів, а доступ до БД обмежений pg_hba.conf: конкретні ролі, підмережа Compose, SCRAM-SHA-256 та reject для решти IPv4/IPv6 адрес.')
 shot('14-security-hba','Права ролей і правила pg_hba.conf')
 shot('15-pgcat-ports','Опублікований лише локальний порт API')
 text('Окремо перевірив docker inspect для всіх контейнерів, відсутність права TRUNCATE у ролі API та відхилення TCP-входу адміністратора правилом pg_hba.conf. Функціональність після обмежень підтверджена Hurl і smoke. Локальний Unix-сокет усередині контейнера залишено trust для адміністративних команд; доступ до Docker тому є адміністративним. TLS між контейнерами в цій локальній лабораторній не налаштовував.')

 heading('8 Дві репліки')
 text('Для logical2 додав postgres_sub2 з томом logical2_data, підпискою tp_sub2 і другим слотом. Обидва subscriber підключаються до publisher незалежно. Перевірки підтвердили 2 000 планів і 6 000 локацій на кожному вузлі, збіг повних контрольних сум та перенесення INSERT, UPDATE і DELETE.')
 shot('16-logical2-equality','Publisher і дві логічні репліки містять однакові дані')
 text('Для PgCat створив дві фізичні репліки replica1 і replica2 з окремими томами й фізичними слотами. Режими two та pgcat використовують synchronous_standby_names=ANY 2 (replica1, replica2) і remote_apply. Обидві копії мають застосувати підтверджений запис; недоступність однієї з них блокує нові commit. Автоматичне перемикання ролі primary не налаштовувалося.')

 heading('9 PgCat і розділення читань та записів')
 text('Додав офіційний образ PgCat 1.2.0 із зафіксованим digest. Шаблон конфігурації містить transaction pooling, query_parser_enabled=true, query_parser_read_write_splitting=true, default_role=primary та primary_reads_enabled=false. Сервери пулу: postgres як primary, replica1 і replica2 як replica. DATABASE_URL додатка тепер вказує на pgcat:6432, реалізація API не змінена.')
 text('Default role primary спрямовує початок явних транзакцій на primary. Звичайні SELECT балансуються між фізичними копіями. Щоб наступний GET після підтвердженого запису бачив актуальні дані незалежно від обраної репліки, обидві копії працюють із remote_apply та ANY 2. Одне асинхронне розділення читань і записів такої гарантії не забезпечує.')
 shot('17-pgcat-routing','Фактичні SELECT потрапляють на дві репліки а UPDATE на primary')
 text('У перевірці маршрутизації виконав 200 SELECT через той самий DATABASE_URL, що використовує API, та UPDATE RETURNING. Усі читання повернули pg_is_in_recovery=true і дві різні адреси; запис повернув false. Це доводить використання обох replica, а не лише факт підключення до PgCat.')
 shot('18-pgcat-hurl','Через PgCat проходять усі 56 оригінальних запитів Hurl')
 sm=load(E/'pgcat-smoke/smoke-summary.json')['metrics']
 text(f'Smoke через PgCat виконав {number(sm["http_reqs"]["values"]["count"],0)} запитів і {number(sm["checks"]["values"]["passes"],0)} успішних перевірок. p95 становить {ms(sm["http_req_duration"]["values"]["p(95)"])}, неочікуваних HTTP помилок {pct(sm["http_req_failed"]["values"]["rate"])}. Усі пороги smoke пройдені.')
 shot('19-pgcat-smoke','Успішний оригінальний smoke k6 через PgCat')
 text('Додатково через PgCat пройшли всі три наявні тести конкурентного доступу: конфлікт версій плану, унікальний порядок паралельно доданих локацій та повторне оновлення після 409. Невдалих тестів немає.')

 heading('10 Порівняння сталого навантаження та читання')
 table(['Топологія','25 VU p95','50 VU p95','100 VU p95','200 VU p95'],[[names[k]]+[number(p[k][str(v)]['p95_ms']) for v in [25,50,100,200]] for k in names],[2.0,1.25,1.25,1.25,1.25])
 text('У таблиці затримки наведені в мілісекундах. Графік показує p95 та кількість завершених HTTP-запитів за секунду на сталих частинах тих самих плато. Підсумки всього запуску з розділу 4 і сталі плато охоплюють різні інтервали.')
 image(E/'load-comparison.png','Затримки й потік запитів на сталих плато однакового Load')
 text('Щоб окремо перевірити читання, використав додатковий сценарій: 200 VU, 45 секунд без пауз, чергування списку планів і деталей свого плану. Для прямого доступу та PgCat використав однакові дві фізичні репліки й ANY 2 remote_apply; змінюється маршрут з API до БД. Виконав по три послідовні запуски, щоразу відновлюючи дані; порівнюю медіани трьох результатів.')
 rows=[]
 for mode,label in [('two','Прямий доступ'),('pgcat','Через PgCat')]:
  v=a['reads'][mode]
  rows.append([label,' / '.join(number(x['rps'],1) for x in v['runs']),number(v['median_rps'],1),ms(v['median_p95_ms'])])
 table(['Маршрут','Запитів за с у трьох запусках','Медіана запитів за с','Медіана p95'],rows,[1.5,2.5,1.5,1.5])
 change=a['read_rps_change_percent']
 text(f'Медіанний потік через PgCat змінився на {number(change,1)} %, медіана p95 - на {number(a["read_p95_change_percent"],1)} %. '+('У цьому сценарії спостерігається приріст швидкодії читання.' if change>0 else 'У цьому сценарії приросту швидкодії читання не отримано.')+' Балансування розподіляє SQL-читання на два процесори БД з окремими лімітами, але додає мережевий перехід, розбір SQL і керування пулом. Підсумок залежить від того, чи є вузьким місцем primary, API, PgCat або спільні ресурси ноутбука; точне профілювання цих складових не виконувалося.')
 shot('20-comparison','Зведені фактичні метрики Load і повторів читання')

 heading('11 Відтворення та висновки')
 text('У README-LAB4.md наведені команди всіх семи топологій. scripts/lab4/run-suite.py запускає їх послідовно, відновлює дані, виконує Hurl, перевірки реплікації, DDL, очікування синхронного commit, smoke, Load та повтори читання. Для повторення потрібно встановити Docker і k6, згенерувати .env та передати шлях до k6 через --k6. Вихідні JSON і текстові журнали збережені в evidence/lab4.')
 text('Логічна реплікація відтворює зміни вибраних таблиць, але схему потрібно підтримувати окремо. Фізична реплікація відтворює кластер і DDL через WAL. Синхронний remote_apply забезпечує видимість підтверджених записів на потрібних копіях та додає очікування й залежність доступності від реплік.')
 text('Без перенаправлення читань додавання репліки саме по собі не зменшує SQL-навантаження primary. PgCat дозволив використати обидві репліки для читань без змін API; функціональні тести та smoke пройшли. Висновки про швидкодію обмежені цим ноутбуком, набором даних, лімітами та виміряними сценаріями.')

 heading('Посилання')
 if pub.get('published'):
  link('Код лабораторної №4',BASE+'/tree/lab4-replication')
  link('Реплікація, безпека та PgCat',BASE+'/commit/'+pub['implementation'])
  link('Конфігурація і сценарії перевірок',BASE+'/commit/'+pub['final_code'])
  if pub.get('results'):link('Вимірювання та скріншоти',BASE+'/commit/'+pub['results'])
 else:
  text('Коміти збережені локально в Git та включені до пакета проєкту.')
 text('На скріншотах показані збережені результати команд; оригінальні файли наведені в evidence/lab4.')
 link('Початкова конфігурація реплікації','https://github.com/potapuff/SumDU.DDS.25F.example/commit/cd66c0f22e08f61602834d4edad4fd82ed672f1a')
 link('Обмеження логічної реплікації PostgreSQL','https://www.postgresql.org/docs/16/logical-replication-restrictions.html')
 link('Режими synchronous_commit та remote_apply','https://www.postgresql.org/docs/16/runtime-config-wal.html')
 link('Параметри PgCat','https://github.com/postgresml/pgcat/blob/main/pgcat.toml')
 # Apply the explicit user typography while preserving original style/package parts.
 for run in d.element.xpath('.//w:r'):
  props=run.find(qn('w:rPr'))
  if props is None:props=OxmlElement('w:rPr');run.insert(0,props)
  font=props.find(qn('w:rFonts'))
  if font is None:font=OxmlElement('w:rFonts');props.insert(0,font)
  font.attrib.clear()
  for key in ['ascii','hAnsi','eastAsia','cs']:font.set(qn('w:'+key),'Times New Roman')
  for tag in ['sz','szCs']:
   size=props.find(qn('w:'+tag))
   if size is None:size=OxmlElement('w:'+tag);props.append(size)
   size.set(qn('w:val'),'28')
 settings=d.settings.element
 if settings.find(qn('w:doNotAutoCompressPictures')) is None:settings.append(OxmlElement('w:doNotAutoCompressPictures'))
 d.core_properties.title='Лабораторна робота 4 Реплікація PostgreSQL у TravelerAPI'
 d.core_properties.subject='Розподілені сховища даних'
 output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);d.save(output)
 with zipfile.ZipFile(REF) as original,zipfile.ZipFile(output) as final:
  preserve=[x for x in original.namelist() if x in ('word/styles.xml','word/stylesWithEffects.xml','word/numbering.xml','word/fontTable.xml') or x.startswith(('word/theme/','word/footer','word/header','customXml/'))]
  temporary=output.with_suffix('.preserve.docx')
  with zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as z:
   for info in final.infolist():z.writestr(info,original.read(info.filename) if info.filename in preserve else final.read(info.filename))
 temporary.replace(output)
 with zipfile.ZipFile(REF) as original,zipfile.ZipFile(output) as final:assert all(original.read(x)==final.read(x) for x in preserve)
 print(json.dumps({'output':str(output),'figures':figure,'preservedParts':len(preserve)},ensure_ascii=False))
if __name__=='__main__':
 import sys;build(sys.argv[1])
